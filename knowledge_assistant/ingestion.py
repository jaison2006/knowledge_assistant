from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

from .core.models import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown", ".html", ".htm", ".csv", ".json"}


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip()


def _document(path: Path, content: str, metadata: dict[str, Any] | None = None) -> Document:
    digest = hashlib.sha256(content.encode("utf-8", errors="ignore")).hexdigest()
    document_metadata = {"filename": path.name, **(metadata or {})}
    return Document(digest[:16], path.name, path.suffix.lower().lstrip(".") or "text", digest, document_metadata)


def ingest_file(path: str | Path, max_bytes: int = 50 * 1024 * 1024) -> tuple[Document, list[dict[str, Any]]]:
    file_path = Path(path).resolve()
    if not file_path.is_file():
        raise FileNotFoundError(file_path)
    if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {file_path.suffix or 'unknown'}")
    if file_path.stat().st_size > max_bytes:
        raise ValueError(f"File exceeds the {max_bytes // (1024 * 1024)} MB limit")

    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _ingest_pdf(file_path)
    if suffix == ".docx":
        return _ingest_docx(file_path)
    if suffix == ".pptx":
        return _ingest_pptx(file_path)
    raw = file_path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        raw = re.sub(r"<[^>]+>", " ", raw)
    elif suffix == ".csv":
        rows = list(csv.reader(io.StringIO(raw)))
        raw = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    elif suffix == ".json":
        raw = json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
    document = _document(file_path, raw)
    return document, [{"content": _clean(raw), "page": None, "section": None, "heading": None}]


def _ingest_pdf(path: Path) -> tuple[Document, list[dict[str, Any]]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF support requires pypdf>=4.0.0") from exc
    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            pages.append({"content": text, "page": number, "section": None, "heading": None})
    document = _document(path, "\n".join(page["content"] for page in pages), {"pages": len(reader.pages)})
    return document, pages


def _ingest_docx(path: Path) -> tuple[Document, list[dict[str, Any]]]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx") from exc
    paragraphs = []
    for paragraph in DocxDocument(str(path)).paragraphs:
        text = _clean(paragraph.text)
        if text:
            heading = text if paragraph.style.name.lower().startswith("heading") else None
            paragraphs.append({"content": text, "page": None, "section": heading, "heading": heading})
    document = _document(path, "\n".join(item["content"] for item in paragraphs))
    return document, paragraphs


def _ingest_pptx(path: Path) -> tuple[Document, list[dict[str, Any]]]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX support requires python-pptx") from exc
    slides = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [_clean(shape.text) for shape in slide.shapes if hasattr(shape, "text") and _clean(shape.text)]
        if texts:
            slides.append({"content": "\n".join(texts), "page": number, "section": texts[0], "heading": texts[0]})
    document = _document(path, "\n".join(item["content"] for item in slides), {"slides": len(slides)})
    return document, slides